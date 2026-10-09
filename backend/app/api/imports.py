import os
import io
import uuid
import datetime
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from typing import Dict, Any, List, Optional
from fastapi import APIRouter, Depends, UploadFile, File, Form, HTTPException, BackgroundTasks
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.models.entities import ImportBatch, SourceRecord, Project
from app.schemas.api_models import (
    UploadPreviewResponse,
    ColumnMapping,
    BatchStartRequest,
    BatchOut,
    SheetPreviewRequest,
    ParsedRecordPreview,
)
from app.crawler.discovery import normalize_url, normalize_domain
from app.services.batch_runner import BatchExecutionRunner

router = APIRouter()

# In-memory storage for raw uploaded file bytes and metadata during import wizard flow
UPLOAD_CACHE: Dict[str, Dict[str, Any]] = {}


def clean_value_for_json(v: Any) -> Any:
    """Recursively converts pandas/numpy/datetime types to JSON-safe primitives."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except Exception:
        pass
    if isinstance(v, (pd.Timestamp, datetime.datetime, datetime.date)):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.isoformat()
    if isinstance(v, (np.integer, int)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        if np.isnan(v) or np.isinf(v):
            return None
        return float(v)
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, dict):
        return {str(k): clean_value_for_json(val) for k, val in v.items()}
    if isinstance(v, (list, tuple, set)):
        return [clean_value_for_json(item) for item in v]
    return str(v) if not isinstance(v, (str, int, float, bool)) else v


def detect_suggested_columns(columns: List[str]) -> ColumnMapping:
    """Intelligently detects common lead spreadsheet column names."""
    cols_clean = [str(c).strip() for c in columns]
    cols_lower = {c.lower(): c for c in cols_clean}

    url_col = ""
    for candidate in ["url", "website", "website url", "company website", "domain", "link", "web", "site"]:
        for k, original in cols_lower.items():
            if candidate == k or candidate in k.split():
                url_col = original
                break
        if url_col:
            break

    name_col = None
    for candidate in ["name", "organization name", "company name", "company", "organization", "account name", "business"]:
        for k, original in cols_lower.items():
            if candidate == k or candidate in k:
                name_col = original
                break
        if name_col:
            break

    industry_col = None
    for candidate in ["industry", "category", "sector", "vertical"]:
        for k, original in cols_lower.items():
            if candidate == k or candidate in k:
                industry_col = original
                break
        if industry_col:
            break

    snov_col = None
    for candidate in ["snov status", "snov result", "snov", "research result", "initial status", "snov_result", "previous status", "qualification", "status"]:
        for k, original in cols_lower.items():
            if candidate == k or candidate in k:
                snov_col = original
                break
        if snov_col:
            break

    location_col = None
    for candidate in ["country", "location", "city", "state", "address"]:
        for k, original in cols_lower.items():
            if candidate == k or candidate in k:
                location_col = original
                break
        if location_col:
            break

    return ColumnMapping(
        company_name_col=name_col,
        url_col=url_col or (cols_clean[0] if cols_clean else ""),
        snov_result_col=snov_col or industry_col,
        industry_col=industry_col,
        location_col=location_col,
    )


def parse_and_preview_sheet(
    filename: str,
    temp_file_id: str,
    content: bytes,
    ext: str,
    sheet_name: Optional[str] = None,
    url_col_override: Optional[str] = None,
    company_name_col_override: Optional[str] = None,
    industry_col_override: Optional[str] = None,
    snov_col_override: Optional[str] = None,
    start_row_override: Optional[int] = None,
    end_row_override: Optional[int] = None,
) -> UploadPreviewResponse:
    """
    Inspects workbook sheets, selects requested sheet (or default 'SNOV SEARCH'),
    and parses strictly that sheet and specified row range.
    """
    # 1. Determine available worksheets
    if ext in (".xlsx", ".xls"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            available_worksheets = list(wb.sheetnames)
            wb.close()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read Excel workbook: {e}")
    else:
        available_worksheets = ["Sheet1"]

    if not available_worksheets:
        raise HTTPException(status_code=400, detail="The uploaded workbook contains no worksheets.")

    # 2. Select worksheet with strict default rule
    if sheet_name:
        if sheet_name not in available_worksheets:
            raise HTTPException(
                status_code=400,
                detail=f"Worksheet '{sheet_name}' does not exist in workbook. Available worksheets: {', '.join(available_worksheets)}"
            )
        selected_worksheet = sheet_name
    else:
        # Default source configuration: check for 'SNOV SEARCH' or sheet with 'snov'
        snov_sheet = next((s for s in available_worksheets if "snov" in s.lower().strip()), None)
        selected_worksheet = snov_sheet if snov_sheet else available_worksheets[0]

    # 3. Read ONLY the selected worksheet
    try:
        if ext == ".csv":
            df = pd.read_csv(io.BytesIO(content))
        else:
            df = pd.read_excel(io.BytesIO(content), sheet_name=selected_worksheet)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read worksheet '{selected_worksheet}': {e}")

    if df.empty:
        raise HTTPException(status_code=400, detail=f"Worksheet '{selected_worksheet}' contains no data rows.")

    columns = [str(c) for c in df.columns]
    column_letters: Dict[str, str] = {}
    for idx, c in enumerate(columns):
        try:
            column_letters[c] = get_column_letter(idx + 1)
        except Exception:
            column_letters[c] = str(idx + 1)

    # 4. Map columns with default for SNOV SEARCH
    if selected_worksheet.strip().lower() == "snov search" or "snov" in selected_worksheet.lower():
        # Default: Column B (url), Column A (name), Column C (industry)
        def_url = next((c for c in columns if c.lower() == "url"), (columns[1] if len(columns) > 1 else columns[0]))
        def_name = next((c for c in columns if c.lower() == "name"), (columns[0] if len(columns) > 0 else None))
        def_ind = next((c for c in columns if c.lower() == "industry"), (columns[2] if len(columns) > 2 else None))

        mapping = ColumnMapping(
            url_col=url_col_override or def_url,
            company_name_col=company_name_col_override if company_name_col_override is not None else def_name,
            industry_col=industry_col_override if industry_col_override is not None else def_ind,
            snov_result_col=snov_col_override if snov_col_override is not None else def_ind,
        )
    else:
        suggested = detect_suggested_columns(columns)
        mapping = ColumnMapping(
            url_col=url_col_override or suggested.url_col,
            company_name_col=company_name_col_override if company_name_col_override is not None else suggested.company_name_col,
            industry_col=industry_col_override if industry_col_override is not None else suggested.industry_col,
            snov_result_col=snov_col_override if snov_col_override is not None else suggested.snov_result_col,
        )

    # Validate url_col if specified
    if mapping.url_col and mapping.url_col not in df.columns:
        raise HTTPException(
            status_code=400,
            detail=f"URL column '{mapping.url_col}' not found in worksheet '{selected_worksheet}'. Available columns: {', '.join(columns)}"
        )

    # 5. Row Range (1-indexed spreadsheet rows, header is row 1)
    total_data_rows = len(df)
    min_data_row = 2
    max_data_row = 1 + total_data_rows

    start_row = start_row_override if (start_row_override and start_row_override >= 2) else min_data_row
    end_row = end_row_override if (end_row_override and end_row_override >= start_row and end_row_override <= max_data_row) else max_data_row

    # Sliced DataFrame for active selection
    start_idx = max(0, start_row - 2)
    end_idx = min(total_data_rows, end_row - 1)
    df_slice = df.iloc[start_idx:end_idx]

    # 6. Analyze URL column strictly within df_slice
    url_col = mapping.url_col
    duplicate_count = 0
    invalid_count = 0
    total_non_empty_urls = 0
    warnings = []

    seen_urls = set()
    if url_col and url_col in df_slice.columns:
        for _, row in df_slice.iterrows():
            val = str(row[url_col]).strip() if pd.notna(row[url_col]) else ""
            if not val or val.lower() in ("nan", "none", "nat"):
                continue
            total_non_empty_urls += 1
            if ("." not in val) or (" " in val and not val.startswith("http")):
                invalid_count += 1
            if val in seen_urls:
                duplicate_count += 1
            else:
                seen_urls.add(val)

        missing_count = len(df_slice) - total_non_empty_urls
        if missing_count > 0:
            warnings.append(f"{missing_count} rows in selected range have missing website URLs")
        if duplicate_count > 0:
            warnings.append(f"Detected {duplicate_count} duplicate URLs in selected range")
        if invalid_count > 0:
            warnings.append(f"Detected {invalid_count} malformed URLs in selected range")
    else:
        warnings.append(f"Please select a valid Website URL column for worksheet '{selected_worksheet}'.")

    # 7. First 20 parsed preview records
    preview_slice = df_slice.head(20)
    preview_records: List[ParsedRecordPreview] = []
    seen_preview_urls = set()

    for idx, (_, row) in enumerate(preview_slice.iterrows()):
        actual_row_num = start_row + idx
        raw_u = str(row[url_col]).strip() if url_col and pd.notna(row.get(url_col)) else ""
        if raw_u.lower() in ("nan", "none", "nat"):
            raw_u = ""
        raw_name = str(row[mapping.company_name_col]).strip() if mapping.company_name_col and pd.notna(row.get(mapping.company_name_col)) else None
        if raw_name and raw_name.lower() in ("nan", "none", "nat"):
            raw_name = None

        is_valid = bool(raw_u and "." in raw_u)
        is_dup = raw_u in seen_preview_urls if raw_u else False
        if raw_u:
            seen_preview_urls.add(raw_u)

        raw_dict = {str(k): clean_value_for_json(v) for k, v in row.items()}
        preview_records.append(
            ParsedRecordPreview(
                row_number=actual_row_num,
                company_name=raw_name,
                url=raw_u,
                is_valid_url=is_valid,
                is_duplicate=is_dup,
                original_row=raw_dict,
            )
        )

    # Legacy preview_rows format for backward compatibility
    preview_rows = [p.original_row for p in preview_records]

    return UploadPreviewResponse(
        filename=filename,
        temp_file_id=temp_file_id,
        available_worksheets=available_worksheets,
        selected_worksheet=selected_worksheet,
        detected_columns=columns,
        column_letters=column_letters,
        suggested_mapping=mapping,
        start_row=start_row,
        end_row=end_row,
        total_detected_rows=len(df_slice),
        total_non_empty_urls=total_non_empty_urls,
        duplicate_url_count=duplicate_count,
        invalid_url_count=invalid_count,
        sample_warnings=warnings,
        preview_rows=preview_rows,
        preview_records=preview_records,
    )


@router.post("/preview", response_model=UploadPreviewResponse)
async def preview_upload(
    file: UploadFile = File(...),
    sheet_name: Optional[str] = Form(None),
):
    """
    Accepts CSV or XLSX, reads worksheet names and schema, and returns preview.
    Caches raw file content for subsequent worksheet/column configuration.
    """
    filename = file.filename or "upload.csv"
    ext = os.path.splitext(filename)[1].lower()

    if ext not in (".csv", ".xlsx", ".xls"):
        raise HTTPException(status_code=400, detail="Only .csv and .xlsx files are supported")

    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File exceeds 50MB limit")

    temp_id = str(uuid.uuid4())

    response = parse_and_preview_sheet(
        filename=filename,
        temp_file_id=temp_id,
        content=content,
        ext=ext,
        sheet_name=sheet_name,
    )

    # Cache in memory
    UPLOAD_CACHE[temp_id] = {
        "filename": filename,
        "content": content,
        "ext": ext,
        "available_worksheets": response.available_worksheets,
        "selected_worksheet": response.selected_worksheet,
    }

    return response


@router.post("/preview-sheet", response_model=UploadPreviewResponse)
async def preview_sheet(payload: SheetPreviewRequest):
    """
    Re-parses an existing uploaded workbook with a different worksheet,
    column selection, or start/end row range without re-uploading the file.
    """
    cached = UPLOAD_CACHE.get(payload.temp_file_id)
    if not cached:
        raise HTTPException(status_code=404, detail="Upload session expired or file not found. Please upload again.")

    response = parse_and_preview_sheet(
        filename=cached["filename"],
        temp_file_id=payload.temp_file_id,
        content=cached["content"],
        ext=cached["ext"],
        sheet_name=payload.sheet_name,
        url_col_override=payload.url_col,
        company_name_col_override=payload.company_name_col,
        industry_col_override=payload.industry_col,
        snov_col_override=payload.snov_result_col,
        start_row_override=payload.start_row,
        end_row_override=payload.end_row,
    )
    cached["selected_worksheet"] = response.selected_worksheet
    return response


@router.post("/start-batch", response_model=BatchOut)
async def start_batch(
    payload: BatchStartRequest,
    background_tasks: BackgroundTasks,
    db: AsyncSession = Depends(get_db)
):
    """
    Confirms worksheet selection and column mapping, creates persistent batch & source records,
    and initiates background crawling & qualification ONLY after explicit confirmation.
    """
    cached = UPLOAD_CACHE.get(payload.temp_file_id)
    if not cached:
        raise HTTPException(status_code=404, detail="Upload session expired or file not found. Please upload again.")

    content = cached["content"]
    ext = cached["ext"]
    filename = payload.filename or cached["filename"]
    available_worksheets = cached.get("available_worksheets", ["Sheet1"])

    project = await db.get(Project, payload.project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # 1. Strictly validate Worksheet existence
    if ext in (".xlsx", ".xls"):
        candidate_sheet = payload.sheet_name or cached.get("selected_worksheet")
        if not candidate_sheet:
            raise HTTPException(
                status_code=400,
                detail=f"Worksheet name must be explicitly specified for Excel files. Available sheets: {', '.join(available_worksheets)}"
            )
        if candidate_sheet not in available_worksheets:
            raise HTTPException(
                status_code=400,
                detail=f"Worksheet '{candidate_sheet}' does not exist in workbook. Available sheets: {', '.join(available_worksheets)}"
            )
        selected_sheet = candidate_sheet
        try:
            df = pd.read_excel(io.BytesIO(content), sheet_name=selected_sheet)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read worksheet '{selected_sheet}': {e}")
    else:
        selected_sheet = payload.sheet_name or "Sheet1"
        df = pd.read_csv(io.BytesIO(content))

    if df.empty:
        raise HTTPException(status_code=400, detail=f"Worksheet '{selected_sheet}' contains no rows.")

    url_col = payload.mapping.url_col
    name_col = payload.mapping.company_name_col
    snov_col = payload.mapping.snov_result_col or payload.mapping.industry_col

    if not url_col or url_col not in df.columns:
        raise HTTPException(
            status_code=400,
            detail=f"Required URL column '{url_col}' not found in worksheet '{selected_sheet}'. Available columns: {', '.join(df.columns)}"
        )

    # 2. Row slicing (1-indexed spreadsheet rows, header is row 1)
    total_data_rows = len(df)
    min_data_row = 2
    max_data_row = 1 + total_data_rows
    start_row = payload.start_row if (payload.start_row and payload.start_row >= 2) else min_data_row
    end_row = payload.end_row if (payload.end_row and payload.end_row >= start_row and payload.end_row <= max_data_row) else max_data_row

    start_idx = max(0, start_row - 2)
    end_idx = min(total_data_rows, end_row - 1)
    df_slice = df.iloc[start_idx:end_idx]

    # Calculate column letter for URL column
    try:
        url_col_idx = list(df.columns).index(url_col)
        url_col_letter = get_column_letter(url_col_idx + 1)
    except Exception:
        url_col_letter = "B"

    total_sliced_rows = len(df_slice)
    source_summary = f"Source: {selected_sheet} | Column {url_col_letter} ({url_col}) | {total_sliced_rows:,} source rows"

    # 3. Create Batch entity
    batch = ImportBatch(
        project_id=payload.project_id,
        filename=filename,
        source_worksheet=selected_sheet,
        source_summary=source_summary,
        start_row=start_row,
        end_row=end_row,
        total_rows=total_sliced_rows,
        status="QUEUED",
        column_mapping={
            **payload.mapping.model_dump(),
            "sheet_name": selected_sheet,
            "start_row": start_row,
            "end_row": end_row,
            "url_col_letter": url_col_letter,
            "source_summary": source_summary,
        },
    )
    db.add(batch)
    await db.flush()

    seen_urls = set()

    for idx, (_, row) in enumerate(df_slice.iterrows()):
        actual_row_num = start_row + idx
        raw_dict = {str(k): clean_value_for_json(v) for k, v in row.items()}
        raw_url = str(row[url_col]).strip() if pd.notna(row.get(url_col)) else ""
        raw_name = str(row[name_col]).strip() if name_col and pd.notna(row.get(name_col)) else None
        raw_snov = str(row[snov_col]).strip() if snov_col and pd.notna(row.get(snov_col)) else None

        if raw_url.lower() in ("nan", "none", "nat"):
            raw_url = ""
        if raw_name and raw_name.lower() in ("nan", "none", "nat"):
            raw_name = None
        if raw_snov and raw_snov.lower() in ("nan", "none", "nat"):
            raw_snov = None

        is_dup = raw_url in seen_urls if raw_url else False
        if raw_url:
            seen_urls.add(raw_url)

        src_record = SourceRecord(
            batch_id=batch.id,
            source_worksheet=selected_sheet,
            row_index=actual_row_num,
            original_data=raw_dict,
            raw_company_name=raw_name,
            raw_url=raw_url,
            raw_snov_result=raw_snov,
            is_duplicate=is_dup,
        )
        db.add(src_record)

    await db.commit()
    await db.refresh(batch)

    # Launch background batch runner ONLY after explicit confirmation
    background_tasks.add_task(BatchExecutionRunner.process_batch, batch.id)

    return BatchOut.model_validate(batch)
